"""Create the synthetic DOCX and PDF samples. Not part of the runtime app."""

from pathlib import Path

import pymupdf
from docx import Document

ROOT = Path(__file__).resolve().parent
DOCUMENTS = ROOT / "documents"

DOCX_TEXT = """APPROVAL_STATUS: sample-not-clinically-approved
DOCUMENT_VERSION: 1.0

SAMPLE ENGINEERING DOCUMENT
This file is synthetic prototype content. It is not a hospital-approved document and it is not for clinical use.

Screening is a clinician decision

Cancer screening means checking people according to a plan set by a qualified clinician or a screening program. This sample states that screening tests are chosen by a clinician. The choice depends on age, personal history, and the clinician's judgment. This prototype does not list ages, brands, or instructions for ordering tests.

A screening result is not a diagnosis. An abnormal screening result needs follow-up with a clinician. A normal screening result does not rule every future illness out, and it does not prove that a person is free of disease. This sample does not interpret a person's result.

Questions a person can ask

People can ask what screening options exist for them, what a result can and cannot show, and when to return for review. This sample does not recommend a particular test, a timetable, or a treatment.
"""

PDF_TEXT = """APPROVAL_STATUS: sample-not-clinically-approved
DOCUMENT_VERSION: 1.0

SAMPLE ENGINEERING DOCUMENT
This file is synthetic prototype content. It is not a hospital-approved document and it is not for clinical use.

Symptoms do not prove or exclude cancer

This sample states a limit on what symptoms mean. A symptom by itself does not prove cancer and does not rule cancer out. Symptoms are changes a person notices, such as pain, fatigue, or a change in weight. The same symptom can have many causes.

Educational use only

Use this sample only to practice document search in an engineering prototype. It must not be used to decide whether someone has cancer or to choose care. Personal concerns belong with a qualified clinician who can take a history and examine the person.

Urgent changes

Sudden severe symptoms need emergency services or urgent in-person care. This document does not assess emergencies and does not replace a call for urgent help.
"""


def write_docx() -> None:
    document = Document()
    lines = DOCX_TEXT.strip().splitlines()
    headings = {
        "SAMPLE ENGINEERING DOCUMENT",
        "Screening is a clinician decision",
        "Questions a person can ask",
    }
    for line in lines:
        text = line.strip()
        if not text:
            continue
        if text in headings:
            document.add_heading(text, level=1)
        else:
            document.add_paragraph(text)
    table = document.add_table(rows=4, cols=2)
    rows = [
        ("Topic", "Sample statement"),
        ("Who chooses a screening test", "A qualified clinician"),
        ("What a screening result is", "Not a diagnosis"),
        ("What this prototype withholds", "No personal test orders and no treatment advice"),
    ]
    for index, (left, right) in enumerate(rows):
        table.rows[index].cells[0].text = left
        table.rows[index].cells[1].text = right
    document.save(DOCUMENTS / "screening_basics.docx")


def write_pdf() -> None:
    document = pymupdf.open()
    page = document.new_page()
    rect = pymupdf.Rect(72, 72, 540, 760)
    unused = page.insert_textbox(rect, PDF_TEXT.strip(), fontsize=11, fontname="helv")
    if unused < 0:
        raise RuntimeError("PDF text did not fit on one page")
    document.save(DOCUMENTS / "symptoms_do_not_prove_cancer.pdf")
    document.close()


if __name__ == "__main__":
    write_docx()
    write_pdf()
    print("wrote sample docx and pdf")
